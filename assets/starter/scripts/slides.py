# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx==1.0.2", "odfpy==1.4.1", "reportlab==4.4.3", "Pillow==11.3.0", "pypdf==6.0.0"]
# ///
"""Export and verify clean keyframes as PPTX, ODP and PDF with speaker notes."""
import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZIP_STORED, ZipFile

from odf import dc, meta
from odf.draw import Frame, Image as OdfImage, Page, PageThumbnail, TextBox
from odf.opendocument import OpenDocumentPresentation
from odf.presentation import Notes
from odf.style import GraphicProperties, MasterPage, PageLayout, PageLayoutProperties, Style, TextProperties
from odf.text import P
from PIL import Image, ImageChops
from pypdf import PdfReader
from pptx import Presentation
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
MIMETYPE = "application/vnd.oasis.opendocument.presentation"
NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "xlink": "http://www.w3.org/1999/xlink",
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sources():
    story_path, timing_path = ROOT / "src/story.json", ROOT / "src/generated/timing.json"
    story = json.loads(story_path.read_text(encoding="utf-8"))
    timing = json.loads(timing_path.read_text(encoding="utf-8"))
    require(re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", story["artifact"]), "Invalid artifact filename")
    require(len(story["scenes"]) == len(timing["scenes"]) > 0, "Scene counts differ")
    require(timing["fps"] == story["fps"] and timing["fps"] > 0, "Frame rates differ")
    seen = set()
    for index, (scene, clock) in enumerate(zip(story["scenes"], timing["scenes"])):
        label = scene["id"]
        require(re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", label), f"Invalid scene ID: {label}")
        require(label not in seen and clock["id"] == label and clock["index"] == index, f"Stale scene: {label}")
        seen.add(label)
        require(isinstance(clock["keyframe"], int) and 0 <= clock["start"] <= clock["keyframe"]
                < clock["start"] + clock["duration"] <= timing["durationInFrames"], f"Invalid keyframe: {label}")
        require(isinstance(scene["narration"], str) and scene["narration"].strip(), f"Missing narration: {label}")
    return story, timing, {"storySha256": digest(story_path), "timingSha256": digest(timing_path)}


def cm(pixels):
    # 144 pixels per physical inch matches the PDF's 0.5-point pixel scale.
    return f"{pixels * 2.54 / 144:.8f}cm"


def make_odp(directory, manifest):
    doc = OpenDocumentPresentation()
    doc.meta.addElement(dc.Title(text=manifest["title"]))
    doc.meta.addElement(meta.Generator(text="Diorama Journey"))
    width, height = cm(manifest["width"]), cm(manifest["height"])
    layout = PageLayout(name="SlideLayout")
    layout.addElement(PageLayoutProperties(pagewidth=width, pageheight=height, margin="0cm", printorientation="landscape"))
    doc.automaticstyles.addElement(layout)
    notes_layout = PageLayout(name="NotesLayout")
    notes_layout.addElement(PageLayoutProperties(pagewidth="21cm", pageheight="29.7cm", margin="0cm", printorientation="portrait"))
    doc.automaticstyles.addElement(notes_layout)
    master = MasterPage(name="Default", pagelayoutname=layout)
    master.addElement(Notes(pagelayoutname=notes_layout))
    doc.masterstyles.addElement(master)
    graphic = Style(name="PlainFrame", family="graphic")
    graphic.addElement(GraphicProperties(stroke="none", fill="none"))
    doc.styles.addElement(graphic)
    paragraph = Style(name="SpeakerText", family="paragraph")
    paragraph.addElement(TextProperties(fontsize="12pt"))
    doc.styles.addElement(paragraph)
    for slide in manifest["slides"]:
        page = Page(name=f"{slide['number']:02}-{slide['id']}", masterpagename=master)
        frame = Frame(name=slide["chapter"], stylename=graphic, x="0cm", y="0cm", width=width, height=height)
        image_path = doc.addPicture(str(directory / slide["image"]))
        frame.addElement(OdfImage(href=image_path, type="simple", show="embed", actuate="onLoad"))
        page.addElement(frame)
        notes = Notes(pagelayoutname=notes_layout)
        notes.addElement(PageThumbnail(pagenumber=slide["number"], **{"class": "page"},
                                      x="2cm", y="1.5cm", width="17cm",
                                      height=f"{17 * manifest['height'] / manifest['width']:.8f}cm"))
        box = Frame(name="Speaker notes", **{"class": "notes"}, placeholder="false", stylename=graphic,
                    x="2cm", y="12cm", width="17cm", height="15cm")
        text = TextBox()
        for line in slide["narration"].split("\n"):
            text.addElement(P(stylename=paragraph, text=line))
        box.addElement(text)
        notes.addElement(box)
        page.addElement(notes)
        doc.presentation.addElement(page)
    doc.save(str(directory / f"{manifest['artifact']}.odp"))


def make_pdf(directory, manifest):
    width, height = manifest["width"] / 2, manifest["height"] / 2
    pdf = Canvas(str(directory / f"{manifest['artifact']}.pdf"), pagesize=(width, height), pageCompression=1)
    pdf.setTitle(manifest["title"])
    pdf.setAuthor("")
    pdf.setCreator("Diorama Journey")
    for slide in manifest["slides"]:
        with Image.open(directory / slide["image"]) as image:
            pdf.drawImage(ImageReader(image.convert("RGB")), 0, 0, width=width, height=height)
        pdf.bookmarkPage(slide["id"])
        pdf.addOutlineEntry(slide["chapter"], slide["id"], level=0)
        pdf.showPage()
    pdf.save()


def make_pptx(directory, manifest):
    presentation = Presentation()
    presentation.slide_width = round(manifest["width"] * 914400 / 144)
    presentation.slide_height = round(manifest["height"] * 914400 / 144)
    presentation.core_properties.title = manifest["title"]
    presentation.core_properties.author = ""
    presentation.core_properties.subject = "Diorama Journey"
    for item in manifest["slides"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(directory / item["image"]), 0, 0,
                                 width=presentation.slide_width, height=presentation.slide_height)
        slide.notes_slide.notes_text_frame.text = item["narration"]
    presentation.save(str(directory / f"{manifest['artifact']}.pptx"))


def verify(directory):
    story, timing, hashes = sources()
    require((directory / "slides.json").is_file(), "Missing slides; run npm run slides")
    manifest = json.loads((directory / "slides.json").read_text(encoding="utf-8"))
    require(manifest["source"] == hashes, "Slides are stale; story or timing changed")
    require(manifest["artifact"] == story["artifact"], "Slide artifact differs from the story")
    require(len(manifest["slides"]) == len(story["scenes"]), "Slide count differs from the story")
    width, height = manifest["width"], manifest["height"]
    require(width > 0 and height > 0 and manifest["captions"] is False, "Invalid slide canvas or subtitles enabled")
    for kind in ("pptx", "odp", "pdf", "notes"):
        item = manifest["files"][kind]
        require(Path(item["file"]).name == item["file"], "Invalid output filename")
        path = directory / item["file"]
        require(path.is_file() and path.stat().st_size == item["bytes"] and digest(path) == item["sha256"], f"Changed or missing {kind} output")
    pdf = PdfReader(directory / manifest["files"]["pdf"]["file"], strict=True)
    require(len(pdf.pages) == len(manifest["slides"]), "PDF page count differs")
    presentation = Presentation(directory / manifest["files"]["pptx"]["file"])
    require(len(presentation.slides) == len(manifest["slides"]), "PPTX page count differs")
    require(presentation.slide_width == round(width * 914400 / 144)
            and presentation.slide_height == round(height * 914400 / 144), "PPTX canvas differs")
    for page, item in zip(presentation.slides, manifest["slides"]):
        require(len(page.shapes) == 1, "PPTX must contain one full-frame image per slide")
        picture = page.shapes[0]
        require(picture.left == picture.top == 0 and picture.width == presentation.slide_width
                and picture.height == presentation.slide_height, "PPTX image is cropped or offset")
        require(hashlib.sha256(picture.image.blob).hexdigest() == item["imageSha256"], "PPTX keyframe differs")
        require(page.notes_slide.notes_text_frame.text == item["narration"], "PPTX speaker notes differ")
    with ZipFile(directory / manifest["files"]["odp"]["file"]) as archive:
        first = archive.infolist()[0]
        require(first.filename == "mimetype" and first.compress_type == ZIP_STORED and not first.extra and not first.comment
                and archive.read("mimetype").decode() == MIMETYPE, "Expected an ODP presentation package")
        content = ET.fromstring(archive.read("content.xml"))
        styles = ET.fromstring(archive.read("styles.xml"))
        layout = styles.find(".//style:page-layout[@style:name='SlideLayout']/style:page-layout-properties", NS)
        require(layout is not None and layout.get(f"{{{NS['fo']}}}page-width") == cm(width)
                and layout.get(f"{{{NS['fo']}}}page-height") == cm(height), "ODP slide size differs")
        pages = content.findall("office:body/office:presentation/draw:page", NS)
        require(len(pages) == len(manifest["slides"]), "ODP page count differs")
        for index, (slide, scene, clock, page, pdf_page) in enumerate(zip(
                manifest["slides"], story["scenes"], timing["scenes"], pages, pdf.pages)):
            label = f"{index + 1:02}-{scene['id']}"
            require(slide["number"] == index + 1 and slide["id"] == scene["id"]
                    and slide["keyframe"] == clock["keyframe"]
                    and slide["narration"] == scene["narration"], f"Wrong slide order, keyframe, or notes: {label}")
            require(slide["image"] == f"frames/{label}.png", f"Unexpected image path: {label}")
            png = directory / slide["image"]
            require(png.is_file() and digest(png) == slide["imageSha256"], f"Missing or changed keyframe: {label}")
            require(page.get(f"{{{NS['draw']}}}name") == label, f"Wrong ODP page order: {label}")
            embedded = page.find("draw:frame/draw:image", NS)
            require(embedded is not None, f"Missing ODP image: {label}")
            image_path = embedded.get(f"{{{NS['xlink']}}}href")
            require(hashlib.sha256(archive.read(image_path)).hexdigest() == slide["imageSha256"], f"Wrong embedded ODP keyframe: {label}")
            paragraphs = page.findall("presentation:notes/draw:frame[@presentation:class='notes']/draw:text-box/text:p", NS)
            notes = "\n".join("".join(p.itertext()) for p in paragraphs)
            require(notes == scene["narration"], f"Missing or altered ODP speaker notes: {label}")
            require(tuple(map(float, pdf_page.mediabox)) == (0, 0, width / 2, height / 2)
                    and pdf_page.rotation == 0, f"Wrong PDF page size: {label}")
            with Image.open(png) as original:
                require(original.size == (width, height), f"Wrong keyframe dimensions: {label}")
                require(len(pdf_page.images) == 1, f"Expected one full-frame PDF image: {label}")
                decoded = pdf_page.images[0].image.convert("RGB")
                require(decoded.size == original.size and ImageChops.difference(original.convert("RGB"), decoded).getbbox() is None,
                        f"PDF keyframe pixels differ: {label}")
    return {"status": "passed", "slides": len(pages), "speakerNotes": len(pages), "width": width, "height": height,
            "keyframes": [s["keyframe"] for s in manifest["slides"]], "files": manifest["files"]}


def build(render_dir):
    require(render_dir.is_relative_to(ROOT / "process/renders"), "Export from this production's render archive")
    story, timing, hashes = sources()
    record = json.loads((render_dir / "render.json").read_text(encoding="utf-8"))
    require(record["mode"] == "slides" and record["status"] == "frames-ready"
            and record["inputProps"]["captions"] is False, "Render clean slide keyframes first")
    require(all(record[key] == value for key, value in hashes.items()), "Story or timing changed during rendering")
    shots = [{"frame": clock["keyframe"], "name": f"{i + 1:02}-{scene['id']}"}
             for i, (scene, clock) in enumerate(zip(story["scenes"], timing["scenes"]))]
    require(record["shots"] == shots, "Render does not contain every chapter keyframe in order")
    directory = render_dir / "slides"
    require(not (directory / "slides.json").exists(), "Do not overwrite a historical slide export")
    manifest = {"schemaVersion": 1, "title": story["title"],
                "artifact": story["artifact"], "sourceCommit": story.get("sourceCommit"),
                "createdAt": datetime.now().astimezone().isoformat(), "source": hashes,
                "renderArchive": render_dir.relative_to(ROOT).as_posix(),
                "width": record["width"], "height": record["height"], "fps": timing["fps"],
                "captions": False, "content": "Full-frame images with editable PPTX and ODP speaker notes", "slides": []}
    for index, (scene, clock, shot) in enumerate(zip(story["scenes"], timing["scenes"], shots)):
        png = directory / "frames" / f"{shot['name']}.png"
        require(png.is_file(), f"Missing keyframe: {png.name}")
        with Image.open(png) as image:
            require(image.size == (record["width"], record["height"]), f"Wrong image size: {png.name}")
            require("A" not in image.getbands() or image.getchannel("A").getextrema() == (255, 255),
                    f"Render an opaque full-frame background: {png.name}")
        manifest["slides"].append({"number": index + 1, "id": scene["id"], "chapter": scene["chapter"],
                                   "title": scene["title"], "keyframe": clock["keyframe"],
                                   "timeSeconds": clock["keyframe"] / timing["fps"],
                                   "narration": scene["narration"], "image": f"frames/{png.name}", "imageSha256": digest(png)})
    make_odp(directory, manifest)
    make_pdf(directory, manifest)
    make_pptx(directory, manifest)
    notes = f"# {story['title'].replace(chr(10), ' ')}\n\n"
    for slide in manifest["slides"]:
        notes += f"## {slide['number']:02} · {slide['chapter']}\n\n{slide['narration']}\n\n"
    (directory / "speaker-notes.md").write_text(notes, encoding="utf-8")
    manifest["files"] = {}
    for kind, name in {"pptx": f"{story['artifact']}.pptx", "odp": f"{story['artifact']}.odp", "pdf": f"{story['artifact']}.pdf", "notes": "speaker-notes.md"}.items():
        path = directory / name
        manifest["files"][kind] = {"file": name, "bytes": path.stat().st_size, "sha256": digest(path)}
    save_json(directory / "slides.json", manifest)
    report = verify(directory)
    save_json(render_dir / "slides-check.json", report)
    # Publish only the verified complete set. Preserve any previous viewing copy.
    public = ROOT / "public"
    with tempfile.TemporaryDirectory(prefix=".slides-", dir=public) as temp:
        staging = Path(temp) / "slides"
        shutil.copytree(directory, staging)
        current, previous = public / "slides", render_dir / "previous-slides"
        if current.exists():
            current.rename(previous)
        try:
            staging.rename(current)
        except BaseException:
            if previous.exists():
                previous.rename(current)
            raise
    output = ROOT / "verification"
    output.mkdir(exist_ok=True)
    save_json(output / f"{datetime.now().astimezone():%Y%m%dT%H%M%S%f%z}-slides-check.json", report)
    sheet = Image.new("RGB", (1280, 360 * ((len(manifest["slides"]) + 1) // 2)), "#f5f2e9")
    for index, item in enumerate(manifest["slides"]):
        with Image.open(directory / item["image"]) as picture:
            sheet.paste(picture.convert("RGB").resize((640, 360), Image.Resampling.LANCZOS),
                        ((index % 2) * 640, (index // 2) * 360))
    (public / "review").mkdir(exist_ok=True)
    sheet.save(public / "review/storyboard.jpg", quality=90)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"Slides: {public / 'slides'}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["build", "verify"])
    parser.add_argument("--render-dir", type=Path, help="Fresh clean-keyframe render; used by render.mjs")
    parser.add_argument("--directory", type=Path, default=ROOT / "public/slides", help="Slide directory to verify")
    args = parser.parse_args()
    if args.command == "build":
        if args.render_dir is None:
            parser.error("build requires --render-dir; normally run npm run slides")
        build(args.render_dir.resolve())
    else:
        print(json.dumps(verify(args.directory.resolve()), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
